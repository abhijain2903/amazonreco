"""ME's trackers on the dashboard: same columns, same order as the team's Excel sheets, built from the hub's records."""
import io
from datetime import timedelta

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from catalog.models import AmazonStock, Forecast, SellOut, Sku
from core import trackers as T
from core.services import vat_h
from orders.models import PurchaseOrder
from promotions.models import Promotion


pytestmark = pytest.mark.django_db

ME_PO = ["PO Number", "Article", "ASIN", "CAT", "WH", "Qty", "Cost", " Delivery value", "Max Hand Off Date", "RFPO", "Sales Order",
         "ASN", "DEL ID", "Del Date", "INVOICE NO", "Inv submission"]
ME_CLAIM = ["Cat.", "VC", "Promotion Title", "MECL Ref. no.", "MECL Ref. no.", "AMZ Agreement ID", "Agr. Status", "Start Date",
            "End Date", "Date Today", "No. of days pending Sub", "MECL Claim submitted", "Submitted Date", "No. of Days for submission",
            "Submitted Value w/o Vat", " Submitted Value w/Vat", " AMZ Claim value"]


def sheet(resp):
    ws = load_workbook(io.BytesIO(resp.content)).active
    return [c.value for c in ws[1]], [[c.value for c in r] for r in ws.iter_rows(min_row=2)]


def test_po_tracker_columns_and_two_booking_portions():
    cols, rows = T.po_tracker(status="all", q="5DRWD3VG")
    assert [c[0] for c in cols[:T.PO_ME]] == ME_PO
    assert [r["cells"][5] for r in rows] == [386, 52]                       # 438 booked in two portions
    so, asn, dl, inv = ({r["cells"][i] for r in rows} for i in (10, 11, 12, 14))
    assert len(so) == len(asn) == len(dl) == len(inv) == 2                   # each portion: its own SO, ASN, delivery, invoice
    assert rows[0]["cells"][7] == 386 * rows[0]["cells"][6]                  # delivery value = qty × cost


def test_po_tracker_shows_what_is_still_to_ship():
    po = PurchaseOrder.objects.get(stage="backorder")
    cols, rows = T.po_tracker(status="all", q=po.po_no)
    open_rows = [r for r in rows if r.get("open")]
    assert open_rows and sum(r["cells"][5] for r in open_rows) == 42 and open_rows[0]["cells"][11] == ""


def test_rfpo_and_references_are_editable(as_user):
    po = PurchaseOrder.objects.filter(stage="booked").first()
    as_user("noura").post(f"/pos/{po.po_no}/references/", {"rfpo": "RFPO-99999"})
    po.refresh_from_db()
    assert po.rfpo == "RFPO-99999"
    p = Promotion.objects.filter(stage="approved").first()
    as_user("reem").post(f"/promos/{p.mecl_ref}/references/", {"sf_ref": "PRO-004559", "brand_ref": "0125CAV-PAMKT158", "subcat": "gpa"})
    p.refresh_from_db()
    assert (p.sf_ref, p.brand_ref, p.subcat) == ("PRO-004559", "0125CAV-PAMKT158", "GPA")


def test_second_portion_gets_its_own_sales_order(as_user):
    po = PurchaseOrder.objects.get(stage="backorder")
    as_user("khalid").post(f"/pos/{po.po_no}/backorder/ship/", {"sales_order": "3009999"})
    po.refresh_from_db()
    assert po.sap_delivery.sales_order == "3009999" and po.sap_delivery.seq == 2


def test_claim_tracker_days_pending_vat_and_amazon_value():
    cols, rows = T.claim_tracker()
    assert [c[0] for c in cols[:T.CLAIM_ME]] == ME_CLAIM
    sub = next(r for r in rows if r["cells"][11] == "Submitted")
    c = sub["cells"]
    assert c[10] == "Done" and c[13] == (c[12] - c[8]).days and c[15] == c[14] + vat_h(c[14])
    assert c[16] is not None                                                  # what Amazon charged (debit notes)
    pend = [r for r in rows if r["cells"][11] == "Pending"]
    assert all(r["cells"][10] == str((timezone.localdate() - r["cells"][8]).days) for r in pend)


def test_sellout_tracker_headers_and_numbers():
    cols, rows, meta = T.sellout_tracker()
    names = [c[0] for c in cols]
    assert names[:6] == ["Model", "ASIN", "Cat", "Status", " RRP", " PO Cost"] and names[6].startswith("SOH ")
    assert names[7] == f"MTD Jan, {meta['year']}" and any(n.endswith("Sellout FCT") for n in names)
    s = Sku.objects.get(model_no=rows[0]["cells"][0])
    jan = sum(SellOut.objects.filter(sku=s, day__year=meta["year"], day__month=1).values_list("units", flat=True))
    assert rows[0]["cells"][7] == jan


def test_u11_and_u12_imports(as_user):
    c = as_user("admin")
    for tid in ("U11", "U12"):
        c.post("/uploads/new/", {"type": tid, "sample": "1"})
        from uploads.models import UploadBatch
        b = UploadBatch.objects.latest("created_at")
        c.post(f"/uploads/{b.pk}/preview/")
        r = c.post(f"/uploads/{b.pk}/commit/")
        assert b"Import complete" in r.content
    assert AmazonStock.objects.exists() and Forecast.objects.exists()


def test_sold_units_come_from_the_sales_report(as_user):
    from promotions.services import stage_of
    p = next(p for p in Promotion.objects.filter(stage="approved") if stage_of(p) in ("live", "waiting_dn") and p.lines.filter(sold_units__isnull=True).exists())
    l = p.lines.filter(sold_units__isnull=True).first()
    SellOut.objects.filter(sku=l.sku).delete()
    SellOut.objects.create(sku=l.sku, day=timezone.localtime(p.start).date() + timedelta(days=1), units=17)
    as_user("reem").post(f"/promos/{p.mecl_ref}/sold/")
    l.refresh_from_db()
    assert l.sold_units == 17


@pytest.mark.parametrize("view,headers", [("po", ME_PO), ("claims", ME_CLAIM), ("sellout", None)])
def test_dashboard_tabs_and_exports(view, headers, as_user):
    c = as_user("tariq")
    r = c.get(f"/?view={view}")
    assert r.status_code == 200 and b"<table" in r.content
    h, data = sheet(c.get(f"/?view={view}&export=xlsx&status=all"))
    if headers:
        assert h == headers
    assert data
    assert c.get("/").status_code == 200                                       # the overview still works
