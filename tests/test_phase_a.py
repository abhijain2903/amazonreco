"""Gap-map phase A: exports, documents, working calendar, dated prices, case packs, master-data edits, credit hold,
disputes with Amazon case IDs and partial recovery, promotion types and several credit notes per claim."""
import io
from datetime import date, datetime, timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import load_workbook

from catalog.models import Price, Sku, agreed_cost_h
from claims.models import Claim
from core import workcal
from core.models import Attachment, Holiday
from orders.models import PurchaseOrder
from orders.services import line_checks
from payments.models import Dispute, Payment
from promotions.models import Promotion

from .conftest import toast

pytestmark = pytest.mark.django_db


# ---------- Excel export ----------
@pytest.mark.parametrize("url", ["/pos/", "/ship/", "/pay/", "/promos/?view=list", "/dns/", "/claims/", "/action/"])
def test_every_list_exports_to_excel(url, as_user):
    r = as_user("admin").get(url + ("&" if "?" in url else "?") + "export=xlsx")
    assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row >= 1 and all(c.value for c in ws[1])


# ---------- documents ----------
def test_attach_pod_to_po_and_view_it(as_user):
    po = PurchaseOrder.objects.filter(delivered_at__isnull=False).first()
    c = as_user("khalid")
    assert b"No proof of delivery attached" in c.get(f"/records/po/{po.po_no}/?tab=shipment").content
    f = SimpleUploadedFile("pod.pdf", b"%PDF-1.4 pod", content_type="application/pdf")
    r = c.post(f"/docs/po/{po.po_no}/", {"file": f, "kind": "pod"})
    assert toast(r)["tone"] != "bad"
    a = Attachment.objects.get(entity="po", entity_id=po.po_no)
    assert a.kind == "pod"
    assert b"pod.pdf" in c.get(f"/records/po/{po.po_no}/?tab=docs").content
    assert b"No proof of delivery attached" not in c.get(f"/records/po/{po.po_no}/?tab=shipment").content
    assert c.get(f"/attachments/{a.pk}/").status_code == 200


def test_attach_needs_upload_permission_and_a_real_record(as_user):
    po = PurchaseOrder.objects.first()
    r = as_user("omar").post(f"/docs/po/{po.po_no}/", {"file": SimpleUploadedFile("x.pdf", b"x"), "kind": "other"})
    assert toast(r)["tone"] == "bad"
    r = as_user("faisal").post("/docs/po/NO-SUCH-PO/", {"file": SimpleUploadedFile("x.pdf", b"x"), "kind": "other"})
    assert r.status_code == 404


def test_generated_files_stay_on_the_record(as_user):
    from core.services import save_file
    po = PurchaseOrder.objects.first()
    f = save_file("po_ack", f"PO_ack_{po.po_no}.csv", [["po_no"], [po.po_no]], "po", po.po_no)
    r = as_user("faisal").get(f"/records/po/{f.entity_id}/?tab=docs")
    assert r.status_code == 200 and f.filename.encode() in r.content


# ---------- working calendar ----------
def test_weekend_and_holidays_are_skipped():
    tz = timezone.get_current_timezone()
    thu = timezone.make_aware(datetime(2026, 10, 1, 16, 0), tz)       # Thursday 16:00
    assert workcal.add_working_hours(thu, 2).date() == date(2026, 10, 4)   # 1 h Thu, then Sunday 09:00
    assert timezone.localtime(workcal.add_working_hours(thu, 2)).hour == 9
    Holiday.objects.create(day=date(2026, 10, 4), name="Test holiday")
    workcal.clear_cache()
    assert timezone.localtime(workcal.add_working_hours(thu, 2)).date() == date(2026, 10, 5)
    assert workcal.add_working_days(thu, 1).date() == date(2026, 10, 5)
    workcal.clear_cache()


def test_fixed_holidays_are_seeded():
    assert Holiday.objects.filter(day=date(2026, 9, 23)).exists() and Holiday.objects.filter(day=date(2026, 2, 22)).exists()


def test_admin_adds_and_removes_a_holiday(as_user):
    c = as_user("admin")
    assert toast(c.post("/settings/holidays/add/", {"day": "2026-12-01", "name": "Test day"}))["tone"] != "bad"
    h = Holiday.objects.get(day=date(2026, 12, 1))
    assert c.post(f"/settings/holidays/{h.pk}/remove/").status_code in (200, 204)
    assert not Holiday.objects.filter(day=date(2026, 12, 1)).exists()
    workcal.clear_cache()


# ---------- prices, case packs, SKU edits ----------
def test_price_valid_on_the_po_date(as_user):
    s = Sku.objects.first()
    old = agreed_cost_h(s, timezone.now())
    tomorrow = (timezone.localdate() + timedelta(days=1)).isoformat()
    r = as_user("admin").post("/settings/prices/save/", {"sku_code": s.sku_code, "cost": (old + 1000) / 100, "valid_from": tomorrow})
    assert toast(r)["tone"] != "bad"
    assert agreed_cost_h(s, timezone.now()) == old                              # today still the old price
    assert agreed_cost_h(s, timezone.now() + timedelta(days=2)) == old + 1000    # the new one from tomorrow
    assert Price.objects.filter(sku=s, valid_to=timezone.localdate()).exists()


def test_case_pack_hint(as_user):
    po = PurchaseOrder.objects.filter(stage="new").first()
    line = po.lines.select_related("sku").first()
    Sku.objects.filter(pk=line.sku_id).update(case_pack=7)
    line.refresh_from_db()
    chk = line_checks(line)
    assert chk["case_pack"] == 7 and chk["case_qty"] % 7 == 0
    assert chk["case_ok"] == (line.qty_ordered % 7 == 0)


def test_sku_add_and_edit(as_user):
    c = as_user("admin")
    data = {"sku_code": "TEST-SKU-1", "model_no": "TST-1", "asin": "B0TEST0001", "category": "PA", "case_pack": "6"}
    assert toast(c.post("/settings/skus/save/", data))["tone"] != "bad"
    assert Sku.objects.get(sku_code="TEST-SKU-1").case_pack == 6
    other = Sku.objects.exclude(sku_code="TEST-SKU-1").first()
    r = c.post("/settings/skus/save/", {**data, "asin": other.asin})
    assert toast(r)["tone"] == "bad" and "already belongs" in toast(r)["msg"]
    assert toast(as_user("faisal").post("/settings/skus/save/", data))["tone"] == "bad"


# ---------- credit hold ----------
def test_credit_hold_needs_a_reason_and_release_lifts_it(as_user):
    po = PurchaseOrder.objects.filter(stage="booked").first()
    c = as_user("omar")
    assert toast(c.post(f"/pos/{po.po_no}/hold/", {"reason": ""}))["tone"] == "bad"
    c.post(f"/pos/{po.po_no}/hold/", {"reason": "Over credit limit"})
    po.refresh_from_db()
    assert po.credit_hold == "Over credit limit" and po.stage == "booked"
    assert b"credit hold" in as_user("admin").get("/action/").content.lower()
    c.post(f"/pos/{po.po_no}/release/")
    po.refresh_from_db()
    assert po.stage == "released" and po.credit_hold == ""


# ---------- short-paid footer ----------
def test_short_paid_po_offers_dispute_not_remittance_upload(as_user):
    p = Payment.objects.filter(status="short", po__isnull=False).first()
    r = as_user("priya").get(f"/records/po/{p.po.po_no}/?tab=invoice")
    assert b"Short-paid by" in r.content and b"Upload the remittance" not in r.content


# ---------- disputes ----------
def test_dispute_case_id_and_partial_recovery(as_user):
    p = Payment.objects.filter(status="short").first()
    c = as_user("priya")
    c.post(f"/pay/{p.payment_no}/dispute/", {"type": "shortage", "amount": p.deduction_h / 100})
    d = Dispute.objects.get(ref=p.payment_no, status="open")
    c.post(f"/pay/disputes/{d.case_no}/submitted/", {"case_id": "VC-123456"})
    d.refresh_from_db()
    assert d.status == "submitted" and d.amazon_case_id == "VC-123456"
    too_much = (d.amount_h + 100) / 100
    assert toast(c.post(f"/pay/disputes/{d.case_no}/won/", {"recovered": too_much}))["tone"] == "bad"
    c.post(f"/pay/disputes/{d.case_no}/won/", {"recovered": d.amount_h / 200})
    d.refresh_from_db()
    assert d.status == "won" and d.recovered_h == d.amount_h // 2
    assert b"Part won" in c.get(f"/records/dispute/{d.case_no}/").content


def test_overdue_dispute_shows_in_action_center(as_user):
    d = Dispute.objects.filter(status__in=["open", "submitted"]).first()
    Dispute.objects.filter(pk=d.pk).update(due=timezone.now() - timedelta(days=1))
    assert f"Follow up dispute {d.case_no}".encode() in as_user("admin").get("/action/").content


# ---------- promotion type ----------
def test_promotion_type_through_the_wizard(as_user):
    c = as_user("reem")
    c.get("/promos/new/")
    c.post("/promos/new/", {"act": "next", "name": "Type test", "cat": "DI", "ptype": "coupon"})
    sku = Sku.objects.filter(category="DI").first()
    c.post("/promos/new/", {"act": "add", "sku": sku.sku_code})
    c.post("/promos/new/", {"act": "next"})
    c.post("/promos/new/", {"act": "draft"})
    p = Promotion.objects.get(name="Type test")
    assert p.promo_type == "coupon"
    assert b"Coupon" in c.get(f"/records/promo/{p.mecl_ref}/").content


# ---------- several credit notes ----------
def test_second_credit_note_closes_a_shortfall(as_user):
    c = Claim.objects.filter(status="sent").first()
    fin = as_user("priya")
    fin.post(f"/claims/{c.claim_no}/cn/", {"cn_no": "CN-A1", "amount": (c.amount_h - 100_00) / 100})
    c.refresh_from_db()
    assert c.status == "shortfall"
    r = fin.post(f"/claims/{c.claim_no}/cn/", {"cn_no": "CN-A1", "amount": "100"})
    assert toast(r)["tone"] == "bad" and "already recorded" in toast(r)["msg"]
    fin.post(f"/claims/{c.claim_no}/cn/", {"cn_no": "CN-A2", "amount": "100"})
    c.refresh_from_db()
    assert c.status == "closed" and c.cn_h == c.amount_h and c.credit_notes.count() == 2
    assert c.promotion.stage == "closed"
    assert toast(fin.post(f"/claims/{c.claim_no}/cn/", {"cn_no": "CN-A3", "amount": "1"}))["tone"] == "bad"


# ---------- go-live guide ----------
def test_uploads_page_has_the_go_live_guide(as_user):
    r = as_user("admin").get("/uploads/")
    assert b"Go-live guide" in r.content and b"of 9 files loaded" in r.content
