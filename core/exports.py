"""Excel export of any list: what is on screen (tab and filters), as an .xlsx download."""
import io
from datetime import date, datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def wants_export(request):
    return request.GET.get("export") == "xlsx"


def sar(h):
    return None if h is None else round(h / 100, 2)


def xlsx(filename, headers, rows):
    """headers: column titles; rows: lists of values (datetimes shown in Riyadh time, money already in SAR)."""
    return xlsx_book(filename, [(filename[:31], headers, rows)])


def xlsx_book(filename, sheets):
    """Several sheets in one workbook: [(title, headers, rows)]."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        _sheet(wb.create_sheet(title[:31]), headers, rows)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}_{timezone.localdate():%Y-%m-%d}.xlsx"'
    return resp


def _sheet(ws, headers, rows):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([timezone.localtime(v).replace(tzinfo=None) if isinstance(v, datetime) and timezone.is_aware(v) else v for v in r])
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None][:500])
        ws.column_dimensions[get_column_letter(i)].width = min(48, width + 2)
        for cell in ws[get_column_letter(i)][1:]:
            if isinstance(cell.value, (datetime, date)):
                cell.number_format = "dd mmm yyyy hh:mm" if isinstance(cell.value, datetime) else "dd mmm yyyy"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
